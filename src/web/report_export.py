# -*- coding: utf-8 -*-
"""报表导出：把 src/web/report.py 的聚合结果写成 Excel（数据 + 说明两个 sheet）。

聚合与参数校验一律复用 report.py，导出路径不重写 SQL 和口径 ——
否则"页面上看到的"和"导出文件里的"会各算各的，这是报表最忌讳的事。
"""

from __future__ import annotations

import io
import logging
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Final, Mapping

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# 让 src/common 与 src/web 能被导入：三种启动方式都能工作
PROJECT_ROOT: Final[Path] = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.common.points import COLUMNS, POINTS   # noqa: E402
from src.web import report   # noqa: E402

# ==================== 1. 配置区（要改参数只动这里） ====================

# type 参数 -> 中文报表名（同时也是文件名里那一段）
REPORT_TYPES: Final[dict[str, str]] = {
    "minute": "分钟报表",
    "day": "日报表",
    "month": "月报表",
    "custom": "自由报表",
}

SHEET_DATA: Final[str] = "数据"
SHEET_INFO: Final[str] = "说明"

HEADER_TIME: Final[str] = "时间"
HEADER_COUNT: Final[str] = "采样条数"
AGGREGATION_NOTE: Final[str] = (
    "各时间窗口内原始采样的算术平均值（TDengine INTERVAL + AVG，聚合在库侧完成）；"
    f"“{HEADER_COUNT}”为该窗口内的原始采样条数"
)
EMPTY_NOTE: Final[str] = "空白 = 该窗口内没有采样（不是 0）；用 0 表示会把“无数据”和“值就是 0”混为一谈"

NUMBER_FORMAT: Final[str] = "0.00"                 # 聚合值保留 2 位小数
TIME_FORMAT: Final[str] = "yyyy-mm-dd hh:mm:ss"
TIME_COLUMN_WIDTH: Final[float] = 20.0
VALUE_COLUMN_WIDTH: Final[float] = 14.0
COUNT_COLUMN_WIDTH: Final[float] = 10.0
INFO_KEY_WIDTH: Final[float] = 14.0
INFO_VALUE_WIDTH: Final[float] = 88.0

# Excel 公式注入防护：这些字符开头的文本会被 Excel 当公式，前面补 ' 变成纯文本
RISKY_PREFIXES: Final[tuple[str, ...]] = ("=", "+", "-", "@")

BOLD_FONT: Final[Font] = Font(bold=True)

# 数据来源写进说明页；与 report.py 用同一份库表配置，避免第三处副本
TD_SOURCE: Final[str] = f"{report.TD_DB}.{report.TD_STABLE}"

# ---- 日志 ----
LOG_LEVEL: Final[int] = logging.INFO
LOG_FORMAT: Final[str] = "%(asctime)s [%(levelname)s] %(name)s: %(message)s"
LOG_DATEFMT: Final[str] = "%Y-%m-%d %H:%M:%S"

LOGGER: Final[logging.Logger] = logging.getLogger("web.report_export")


class ExportParamError(ValueError):
    """导出参数不合法（type 缺失/非法）—— 接口层转 400。"""


# ==================== 2. 小工具 ====================

def as_datetime(text: Any) -> Any:
    """把 'YYYY-MM-DD HH:MM:SS' 转成 datetime（Excel 里可排序可筛选）；转不动就原样写文本。"""
    try:
        return datetime.strptime(str(text), report.TS_FORMAT)
    except ValueError:
        return as_text(text)


def as_number(value: Any) -> Any:
    """数值保留 2 位小数；None 保持 None —— 在 Excel 里就是空白，绝不能写成 0。"""
    if isinstance(value, (int, float)):
        return round(float(value), report.ROUND_DIGITS)
    return None


def as_text(value: Any) -> str:
    """文本单元格：挡住以 = + - @ 开头的值，避免被 Excel 当公式执行。"""
    text = "" if value is None else str(value)
    return f"'{text}" if text[:1] in RISKY_PREFIXES else text


def build_filename(report_type: str, envelope: dict[str, Any]) -> str:
    """生成中文文件名，例如 CEMS_日报表_2026-09-11.xlsx。"""
    start = str(envelope["start"])
    if report_type == "day":
        stamp = start[:10]                                   # 2026-09-11
    elif report_type == "month":
        stamp = start[:7]                                    # 2026-09
    else:
        stamp = start[:16].replace(" ", "_").replace(":", "")  # 2026-09-11_2229
    return f"CEMS_{REPORT_TYPES[report_type]}_{stamp}.xlsx"


# ==================== 3. 工作簿组装 ====================

def run_report(report_type: str, query: Mapping[str, str], scope: str = "") -> dict[str, Any]:
    """按 type 调 report.py 里对应的报表函数（复用同一套校验与聚合口径）。"""
    if report_type == "minute":
        return report.minute_report(query.get("start"), query.get("end"), scope)
    if report_type == "day":
        return report.day_report(query.get("date"), scope)
    if report_type == "month":
        return report.month_report(query.get("year"), query.get("month"), scope)
    return report.custom_report(query.get("start"), query.get("end"), scope)


def fill_data_sheet(sheet: Any, envelope: dict[str, Any]) -> None:
    """数据 sheet：A 列时间，随后 8 个测点列，最后采样条数。

    表头列名与单位直接取自 src/common/points.py，不再手抄第 N 份契约。
    """
    sheet.column_dimensions["A"].width = TIME_COLUMN_WIDTH
    for index in range(2, len(POINTS) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = VALUE_COLUMN_WIDTH
    sheet.column_dimensions[get_column_letter(len(POINTS) + 2)].width = COUNT_COLUMN_WIDTH
    sheet.freeze_panes = "A2"          # 表头冻结，翻长报表时不会看丢列名

    header = [WriteOnlyCell(sheet, value=HEADER_TIME)]
    header += [
        WriteOnlyCell(sheet, value=as_text(f"{point.name} ({point.unit})")) for point in POINTS
    ]
    header.append(WriteOnlyCell(sheet, value=HEADER_COUNT))
    for cell in header:
        cell.font = BOLD_FONT
    sheet.append(header)

    for point in envelope["points"]:
        time_cell = WriteOnlyCell(sheet, value=as_datetime(point["ts"]))
        time_cell.number_format = TIME_FORMAT
        cells = [time_cell]
        for column in COLUMNS:
            value_cell = WriteOnlyCell(sheet, value=as_number(point.get(column)))
            value_cell.number_format = NUMBER_FORMAT
            cells.append(value_cell)
        cells.append(WriteOnlyCell(sheet, value=point.get("n")))
        sheet.append(cells)


def fill_info_sheet(sheet: Any, report_type: str, envelope: dict[str, Any]) -> None:
    """说明 sheet：报表类型/口径/时间范围/生成时间/数据来源/空值含义。

    报表拿去存档或上报时，这页是刚需 —— 否则没人知道数字是怎么算出来的。
    """
    sheet.column_dimensions["A"].width = INFO_KEY_WIDTH
    sheet.column_dimensions["B"].width = INFO_VALUE_WIDTH
    rows = (
        ("报表类型", REPORT_TYPES[report_type]),
        ("统计口径", AGGREGATION_NOTE),
        ("时间范围", f"{envelope['start']} ~ {envelope['end']}（窗口单位 {envelope['unit']}）"),
        ("窗口数", str(len(envelope["points"]))),
        ("生成时间", datetime.now().strftime(report.TS_FORMAT)),
        ("数据来源", f"{TD_SOURCE}（TDengine）"),
        ("空值含义", EMPTY_NOTE),
    )
    for key, value in rows:
        key_cell = WriteOnlyCell(sheet, value=as_text(key))
        key_cell.font = BOLD_FONT
        sheet.append([key_cell, WriteOnlyCell(sheet, value=as_text(value))])


def build_workbook(report_type: str, envelope: dict[str, Any]) -> bytes:
    """生成 xlsx 字节：write_only 流式写，回报表点数上限下内存占用很小。"""
    workbook = Workbook(write_only=True)
    fill_data_sheet(workbook.create_sheet(SHEET_DATA), envelope)
    fill_info_sheet(workbook.create_sheet(SHEET_INFO), report_type, envelope)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_workbook(query: Mapping[str, str], scope: str = "") -> tuple[bytes, str]:
    """导出入口：返回 (xlsx 字节, 中文文件名)。

    参数非法抛 ExportParamError / report.ReportParamError，查库失败抛 report.ReportQueryError，
    由接口层分别转 400 与 503（错误不会写进 xlsx 里让用户下载）。
    """
    report_type = (query.get("type") or "").strip()
    if report_type not in REPORT_TYPES:
        raise ExportParamError(
            f"type 应为 {'/'.join(REPORT_TYPES)} 之一，收到 {report_type!r}"
        )
    envelope = run_report(report_type, query, scope)
    content = build_workbook(report_type, envelope)
    filename = build_filename(report_type, envelope)
    LOGGER.info(
        "导出 %s：%d 个窗口，%d 字节，文件名 %s",
        REPORT_TYPES[report_type], len(envelope["points"]), len(content), filename,
    )
    return content, filename
