# -*- coding: utf-8 -*-
"""验收自检：报表接口的"覆盖率 + 样本不足"标记（只读，走 HTTP，不改任何数据）。

被验的对象是**拔线演练暴露的真实缺陷**：数据缺口在报表里是"消失"的，而不是"被标记"的。
本脚本用库里 2026-10-01 21:26 那一段**真实断档**做回归，逐条断言：

  A. 断档窗口必须**出现在结果里**：21:26 ⇒ n=0、coverage=0.0、insufficient=true
  B. 断档两侧：21:24（11 条）coverage=0.9167 不判不足；
     21:27（9 条）coverage=0.75 **正好等于门限**⇒ 按"覆盖率 < 门限"的口径不判不足
     （与 docs/告警判据设计.md §3.4 的 `覆盖率 < 门限` 一致：等于门限不算不足）
  C. 正常窗口不被误标：21:28（12 条）coverage=1.0、insufficient=false
  D. envelope 汇总自洽：窗口数/不足窗口数/整体覆盖率/缺口清单互相对得上，
     且"逐点 insufficient"与"汇总 insufficient_windows"完全一致（不会各说各话）
  E. 三种粒度 1m/1h/1d 各跑一次，覆盖率与判定都拿到（并打印数值）
  F. 契约只增不减：unit/start/end/points 与 point 的 ts/9 个测点/n 全部在位，无改名
  G. Excel 导出列布局未动：表头仍是"时间 + 9 测点 + 采样条数"，行数等于窗口数
  H. 曲线接口未被破坏（聚合粒度带上同一份 coverage 汇总）

用法（需要 8 个服务已在跑）：
    python scripts/verify_report_coverage.py
    python scripts/verify_report_coverage.py --base-url http://127.0.0.1:5001
    python scripts/verify_report_coverage.py --gap-start "2026-10-01 21:24" --gap-end "2026-10-01 21:29"

设备维度（这段的 SQL 必须按设备过滤）
-----------------------------------------------------------------------------
`cems_data` 是多设备共用的超级表（TAG = plant/device），两台设备在同时写。
下面那段**前置核查 SQL 必须带 `device = '<scope>'`**：不带过滤时同一分钟的
COUNT(*) 会把两台设备（外加压测标签）的行一起算上，断档分钟被另一台的数据填满，
"确认 21:26 那一分钟不返回"的核查结论会直接反过来
（实测 2026-10-02 最近 60 分钟同一段聚合：不过滤合计 1372 行，按 device1 = 714 行、
device2 = 658 行，混读正好是两台之和）。

⚠️ 本脚本验的是**展示层接口**，而接口的查询在服务端就已经按设备过滤了
（`src/web/report.py` 从 `cache.DEVICE_SCOPE_PARTS` 取 plant/device，
值来自 web 容器的 `TD_PLANT`/`TD_DEVICE`）。所以这里**没有**、也不该有
`AND device = ...` —— HTTP 层压根没有设备参数，加了只会制造"以为按设备查了"的假象。
替代做法：用 `--device` 声明"我期望验的是哪台设备"，脚本会读 web 实例
`/api/cache/stats` 的 `device_scope` 核对，**不一致就直接失败**，
避免拿着 device1 的接口结果去证明 device2 的覆盖率。

前置：断档区间必须真的在**被测设备**的库里（先跑一次
      docker exec tdengine taos -s "SELECT _wstart, COUNT(*) FROM cems.cems_data
        WHERE ts >= '2026-10-01 21:24:00' AND ts < '2026-10-01 21:29:00'
          AND device = 'device1' INTERVAL(1m)"
      确认 21:26 那一分钟不返回）。
"""

from __future__ import annotations

import argparse
import io
import json
import os
import sys
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any, Final

PROJECT_ROOT: Final[Path] = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
if str(PROJECT_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

# 期望的表头真源取自导出模块本身，脚本里不再抄第三份列名
from src.common.points import COLUMNS, POINTS          # noqa: E402
from src.web import report_export                      # noqa: E402

from _device_scope import add_device_argument, resolve_device   # noqa: E402

DEFAULT_BASE_URL: Final[str] = "http://127.0.0.1"          # 对外入口（nginx:80）
DIRECT_BASE_URL: Final[str] = "http://127.0.0.1:5001"      # 直连 web 容器端口（对照）
DEFAULT_GAP_START: Final[str] = "2026-10-01 21:24"
DEFAULT_GAP_END: Final[str] = "2026-10-01 21:29"
TIMEOUT_SECONDS: Final[float] = 20.0
#: 直连 web 容器读它的设备维度（nginx 不转发 /api/cache/stats，实测 403）
STATS_BASE_URL: Final[str] = os.getenv("WEB_DIRECT_BASE_URL", DIRECT_BASE_URL)

FAILURES: list[str] = []


def _check(condition: bool, description: str) -> None:
    """记录一条验收判定；失败时收进 FAILURES，最后统一汇总。"""
    print(f"  [{'PASS' if condition else 'FAIL'}] {description}")
    if not condition:
        FAILURES.append(description)


def get_json(base_url: str, path: str, params: dict[str, str]) -> dict[str, Any]:
    """GET 一个 JSON 接口；HTTP 非 200 或响应不是 JSON 一律当验收失败抛出。"""
    from urllib.parse import urlencode
    url = f"{base_url.rstrip('/')}{path}?{urlencode(params)}"
    try:
        with urllib.request.urlopen(url, timeout=TIMEOUT_SECONDS) as response:
            body = response.read().decode("utf-8")
    except (urllib.error.URLError, OSError) as exc:
        raise SystemExit(f"连不上 {url}: {exc}（确认 nginx / web 容器在跑）") from exc
    return json.loads(body)


def get_bytes(base_url: str, path: str, params: dict[str, str]) -> tuple[bytes, dict[str, str]]:
    """GET 一个二进制接口（Excel 导出），返回 (内容, 响应头)。"""
    from urllib.parse import urlencode
    url = f"{base_url.rstrip('/')}{path}?{urlencode(params)}"
    try:
        with urllib.request.urlopen(url, timeout=TIMEOUT_SECONDS) as response:
            return response.read(), dict(response.headers)
    except (urllib.error.URLError, OSError) as exc:
        raise SystemExit(f"连不上 {url}: {exc}") from exc


def by_ts(envelope: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {point["ts"]: point for point in envelope["points"]}


def _minute(start_text: str, offset: int) -> str:
    """把 'YYYY-MM-DD HH:MM' 往后推 `offset` 分钟，返回 'YYYY-MM-DD HH:MM:00'。

    断档区间由命令行给出，脚本不硬编码具体分钟 —— 换一段断档重跑不用改代码。
    """
    from datetime import datetime, timedelta
    moment = datetime.strptime(start_text[:16], "%Y-%m-%d %H:%M")
    return (moment + timedelta(minutes=offset)).strftime("%Y-%m-%d %H:%M:00")


def section_0_device_scope(device: str) -> None:
    """0. 设备维度自检：接口服务端实际查的是哪台设备，必须与 --device 一致。

    多设备上线后，`cems_data` 里两台设备同时在写，展示层的查询按
    `TD_PLANT`/`TD_DEVICE`（⇒ `cache.DEVICE_SCOPE`）过滤。如果只看 HTTP 结果、
    不核对服务端 scope，就可能拿 device1 的接口结果去声称"device2 覆盖率达标"。

    读 `/api/cache/stats` 的 `device_scope`（直连 web 容器：nginx 不转发该路径）。
    取不到就**跳过**而不是判失败 —— 直接连 5001 的场景是可选的，
    缺这条元数据不代表覆盖率断言本身失效（诚实写明，不假装验过）。
    """
    print("\n=== 0. 设备维度自检（HTTP 接口的服务端 scope）===")
    url = f"{STATS_BASE_URL.rstrip('/')}/api/cache/stats"
    try:
        with urllib.request.urlopen(url, timeout=TIMEOUT_SECONDS) as response:
            stats = json.loads(response.read().decode("utf-8"))
    except (urllib.error.URLError, OSError, ValueError) as exc:
        print(f"  [SKIP] 读不到 {url}（{exc}）；"
              f"无法核对服务端设备维度，请自行确认接口确为 {device}")
        return
    scope = str(stats.get("device_scope", ""))
    print(f"  期望设备={device}；web 实例 device_scope={scope or '(未下发)'}")
    _check(scope.endswith(f"/{device}") or scope == device,
           f"0. web 实例的 device_scope({scope}) 与 --device({device}) 一致")


def section_a_gap_visible(base_url: str, start: str, end: str) -> dict[str, Any]:
    """A/B/C. 断档窗口必须出现且被标记；两侧与正常窗口的判定必须正确。"""
    print(f"\n=== A/B/C. 真实断档区间的逐窗标记（{start} ~ {end}，粒度 1m）===")
    envelope = get_json(base_url, "/api/report/minute", {"start": start, "end": end})
    points = by_ts(envelope)
    print(f"{'窗口':<21}{'条数':>5}{'应有':>6}{'覆盖率':>9}{'不足':>7}{'未到期':>8}")
    for point in envelope["points"]:
        print(f"{point['ts']:<21}{point['n']:>5}{point['expected']:>6}"
              f"{point['coverage']:>9}{str(point['insufficient']):>7}{str(point['pending']):>8}")

    gap_ts = _minute(start, 2)             # 实测断档：21:26 整分钟无数据
    before_ts = _minute(start, 0)          # 断档前：21:24（11 条）
    after_ts = _minute(start, 3)           # 断档后：21:27（9 条）
    normal_ts = _minute(start, 4)          # 恢复正常：21:28（12 条）

    _check(gap_ts in points, f"A. 断档窗口 {gap_ts} 出现在结果里（不再消失）")
    if gap_ts in points:
        gap = points[gap_ts]
        _check(gap["n"] == 0, f"A. {gap_ts} n == 0（整窗无数据）")
        _check(gap["coverage"] == 0.0, f"A. {gap_ts} coverage == 0.0")
        _check(gap["insufficient"] is True, f"A. {gap_ts} insufficient == true")
        _check(not gap["pending"], f"A. {gap_ts} 不是 pending（它是过去的真缺口，不是未到期）")
        _check(all(gap[column] is None for column in COLUMNS),
               f"A. {gap_ts} 9 个测点全是 null（不许用 0 或推算值填）")

    if before_ts in points:
        before = points[before_ts]
        _check(before["n"] == 11, f"B. {before_ts} n == 11（实测）")
        _check(before["coverage"] == 0.9167, f"B. {before_ts} coverage == 0.9167")
        _check(before["insufficient"] is False, f"B. {before_ts} 不判样本不足")
    else:
        _check(False, f"B. {before_ts} 未出现在结果里")

    if after_ts in points:
        after = points[after_ts]
        _check(after["n"] == 9, f"B. {after_ts} n == 9（实测）")
        _check(after["coverage"] == 0.75, f"B. {after_ts} coverage == 0.75（正好等于门限）")
        _check(after["insufficient"] is False,
               f"B. {after_ts} 等于门限不算不足（口径：覆盖率 < 门限才判不足）")
    else:
        _check(False, f"B. {after_ts} 未出现在结果里")

    if normal_ts in points:
        normal = points[normal_ts]
        _check(normal["n"] == 12, f"C. {normal_ts} n == 12（正常窗口）")
        _check(normal["coverage"] == 1.0, f"C. {normal_ts} coverage == 1.0")
        _check(normal["insufficient"] is False, f"C. 正常窗口 {normal_ts} 未被误标为样本不足")
    else:
        _check(False, f"C. {normal_ts} 未出现在结果里")
    return envelope


def section_d_summary(envelope: dict[str, Any]) -> None:
    """D. 汇总与逐点标记必须自洽（不会各说各话）。"""
    print("\n=== D. envelope 汇总自洽性 ===")
    coverage = envelope.get("coverage")
    if not isinstance(coverage, dict):
        _check(False, "D. 响应里缺少 coverage 汇总")
        return
    print("  汇总: " + json.dumps(coverage, ensure_ascii=False))
    threshold = coverage["threshold"]
    windows = coverage["windows"]
    flagged = [point for point in envelope["points"] if point["insufficient"]]
    _check(windows == len(envelope["points"]), "D. coverage.windows == len(points)")
    _check(coverage["insufficient_windows"] == len(flagged),
           "D. coverage.insufficient_windows == 逐点 insufficient 的个数")
    _check(coverage["gap_count"] == len(flagged) or coverage["gaps_truncated"],
           "D. gap_count 与不足窗口数一致（未截断时）")
    _check(coverage["threshold"] == threshold and threshold > 0, "D. 门限随响应一起下发")
    judged = [point for point in envelope["points"] if not point["pending"]]
    _check(coverage["judged_windows"] == len(judged), "D. judged_windows == 非 pending 的窗口数")
    _check(coverage["pending_windows"] == len(envelope["points"]) - len(judged),
           "D. pending_windows == pending 的窗口数")
    # 整体覆盖率 = 实际条数 / 应有条数（两边都按各窗口自己的 expected 求和）
    actual = sum(point["n"] for point in envelope["points"])
    expected = sum(point["expected"] for point in envelope["points"])
    _check(coverage["actual_total"] == actual and coverage["expected_total"] == expected,
           "D. actual_total / expected_total 与逐窗求和一致")
    if expected:
        _check(abs(coverage["overall"] - min(1.0, actual / expected)) < 1e-4,
               "D. overall == 实际条数 / 应有条数")
    # 缺口清单每一行都能复算
    for gap in coverage["gaps"]:
        point = by_ts(envelope)[gap["ts"]]
        if not (gap["n"] == point["n"] and gap["missing"] == point["expected"] - point["n"]
                and gap["coverage"] < threshold):
            _check(False, f"D. 缺口清单行 {gap['ts']} 与 points 对不上")
            return
    _check(all(point["coverage"] >= threshold for point in envelope["points"]
               if not point["pending"] and not point["insufficient"]),
           "D. 不存在'未判不足但覆盖率低于门限'的窗口")


def section_e_three_units(base_url: str, gap_start: str, gap_end: str) -> None:
    """E. 三种窗口粒度各跑一次，覆盖率与判定都要拿得到。"""
    print("\n=== E. 三种窗口粒度各跑一次 ===")
    day = gap_start[:10]
    cases: list[tuple[str, str, dict[str, str]]] = [
        ("1m", "/api/report/minute", {"start": f"{day} 19:00", "end": gap_end}),
        ("1h", "/api/report/custom", {"start": f"{day} 00:00", "end": f"{day} 23:00"}),
        ("1d", "/api/report/month", {"year": day[:4], "month": str(int(day[5:7]))}),
    ]
    for unit, path, params in cases:
        envelope = get_json(base_url, path, params)
        coverage = envelope.get("coverage")
        if not isinstance(coverage, dict):
            _check(False, f"E. [{unit}] 响应里缺少 coverage 汇总")
            continue
        print(f"  [{unit}] unit={envelope['unit']} 窗口={coverage['windows']} "
              f"未到期={coverage['pending_windows']} 不足={coverage['insufficient_windows']} "
              f"整体覆盖率={coverage['overall']} 实际={coverage['actual_total']} "
              f"应有={coverage['expected_total']}（{envelope['start']} ~ {envelope['end']}）")
        _check(envelope["unit"] == unit, f"E. [{unit}] 返回的 unit 是 {unit}")
        _check(coverage["windows"] > 0, f"E. [{unit}] 窗口数 > 0")
        _check(coverage["overall"] is not None, f"E. [{unit}] 有整体覆盖率数值")
        _check(0.0 <= coverage["overall"] <= 1.0, f"E. [{unit}] 整体覆盖率落在 0~1")
        _check(all(0.0 <= point["coverage"] <= 1.0 for point in envelope["points"]),
               f"E. [{unit}] 逐窗覆盖率都落在 0~1")


def section_f_contract(base_url: str, start: str, end: str) -> None:
    """F. 既有字段只能加、不能删或改名。"""
    print("\n=== F. 契约只增不减 ===")
    envelope = get_json(base_url, "/api/report/minute", {"start": start, "end": end})
    for key in ("unit", "start", "end", "points"):
        _check(key in envelope, f"F. envelope 保留字段 {key}")
    _check(isinstance(envelope["points"], list) and envelope["points"], "F. points 仍是数组且非空")
    point = envelope["points"][0]
    for key in ("ts", "n", *COLUMNS):
        _check(key in point, f"F. point 保留字段 {key}")
    for key in ("coverage", "insufficient", "expected", "pending"):
        _check(key in point, f"F. point 新增字段 {key}")


def section_g_export(base_url: str, start: str, end: str) -> None:
    """G. Excel 导出的列布局未动（前端/Excel 属"押后"那批，本轮不能碰）。"""
    print("\n=== G. Excel 导出列布局未变 ===")
    content, headers = get_bytes(base_url, "/api/report/export",
                                 {"type": "minute", "start": start, "end": end})
    _check(content[:2] == b"PK", "G. 导出内容是 xlsx（ZIP 头）")
    _check("spreadsheetml" in headers.get("Content-Type", ""), "G. Content-Type 是 xlsx")
    try:
        from openpyxl import load_workbook
    except ImportError:
        _check(False, "G. 本机没有 openpyxl，无法核对列布局")
        return
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    expected_header = (
        [report_export.HEADER_TIME]
        + [f"{point.name} ({point.unit})" for point in POINTS]
        + [report_export.HEADER_COUNT]
    )
    print(f"  表头: {header}")
    _check(header == expected_header, "G. 数据表表头与改动前一致（时间 + 9 测点 + 采样条数）")
    rows = sum(1 for _ in sheet.iter_rows(min_row=2))
    envelope = get_json(base_url, "/api/report/minute", {"start": start, "end": end})
    _check(rows == len(envelope["points"]),
           f"G. 数据行数 {rows} == 接口窗口数 {len(envelope['points'])}")


def section_h_curve(base_url: str, gap_start: str, gap_end: str) -> None:
    """H. 曲线接口（复用 report.py 的口径）未被破坏。"""
    print("\n=== H. 曲线接口 ===")
    curve = get_json(base_url, "/api/curve", {"start": f"{gap_start[:10]} 19:00", "end": gap_end})
    _check(all(key in curve for key in ("unit", "ts", "points", "start", "end")),
           "H. 曲线保留字段 unit/ts/points/start/end")
    _check(len(curve["ts"]) == curve["points"], "H. 列式 ts 长度 == points")
    _check(isinstance(curve.get("coverage"), dict),
           "H. 聚合粒度曲线带上覆盖率汇总（与报表同一份）")
    raw = get_json(base_url, "/api/curve",
                   {"start": f"{gap_end[:10]} {gap_end[11:13]}:44", "end": f"{gap_end[:10]} {gap_end[11:13]}:45"})
    _check(raw["unit"] == "raw" and "coverage" not in raw,
           "H. 原始点粒度不带 coverage（没有窗口这个概念，不硬塞占位字段）")


def main() -> int:
    parser = argparse.ArgumentParser(description="验收报表覆盖率与样本不足标记")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL,
                        help=f"接口入口（默认 nginx: {DEFAULT_BASE_URL}；直连: {DIRECT_BASE_URL}）")
    parser.add_argument("--gap-start", default=DEFAULT_GAP_START, help="断档区间起点")
    parser.add_argument("--gap-end", default=DEFAULT_GAP_END, help="断档区间终点")
    add_device_argument(
        parser,
        help_text=(
            "期望验的是哪台设备（默认 device1）。"
            "本脚本走 HTTP，查询由展示层按它自己的 TD_PLANT/TD_DEVICE 过滤，"
            "这里用来自检二者一致（不一致直接失败）"
        ),
    )
    args = parser.parse_args()

    device = resolve_device(args.device)
    print(
        f"报表覆盖率验收：base_url={args.base_url} 断档区间={args.gap_start} ~ {args.gap_end} "
        f"设备={device}（HTTP 侧由展示层按设备过滤，本脚本核对一致）"
    )
    section_0_device_scope(device)
    envelope = section_a_gap_visible(args.base_url, args.gap_start, args.gap_end)
    section_d_summary(envelope)
    section_e_three_units(args.base_url, args.gap_start, args.gap_end)
    section_f_contract(args.base_url, args.gap_start, args.gap_end)
    section_g_export(args.base_url, args.gap_start, args.gap_end)
    section_h_curve(args.base_url, args.gap_start, args.gap_end)

    print("\n================ 结论 ================")
    if FAILURES:
        print(f"不通过 {len(FAILURES)} 项:")
        for item in FAILURES:
            print(f"  - {item}")
        return 1
    print("全部通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
